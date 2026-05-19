import os
import PyPDF2
from docx import Document

def read_docx(path):
    doc = Document(path)
    res = []
    for para in doc.paragraphs:
        if para.text.strip():
            res.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip().replace('\n', ' '))
            res.append(" | ".join(row_data))
    return "\n".join(res)

def read_pdf(path):
    res = []
    with open(path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            res.append(page.extract_text())
    return "\n".join(res)

print("=== KPA.docx ===")
print(read_docx('KPA.docx')[:1000])

print("\n=== CEP-01.docx.pdf ===")
print(read_pdf('CEP-01.docx.pdf')[:1000])

print("\n=== Justification Form of project.docx ===")
print(read_docx('Justification Form of project.docx'))
